import os
from datetime import datetime
from pathlib import Path
from typing import Any

from wifi_risk.services.assessment_service import get_assessment_by_id
from wifi_risk.services.device_service import get_device_by_id
from wifi_risk.services.finding_service import (
    count_findings_by_severity,
    get_findings_by_assessment_id,
    get_findings_by_device_id,
)
from wifi_risk.services.scoring_service import evaluate_device_security

REPORTS_DIR = "reports"


def get_assessment_report_data(assessment_id: str) -> dict[str, Any]:
    """
    Compile all data required for generating comprehensive assessment reports.
    """
    asm = get_assessment_by_id(assessment_id)
    if not asm:
        raise ValueError(f"Assessment ID not found: {assessment_id}")

    device_id = asm.get("device_id", "WR-001")
    dev = get_device_by_id(device_id)

    findings = get_findings_by_assessment_id(assessment_id)
    if not findings:
        findings = get_findings_by_device_id(device_id)

    evaluation = evaluate_device_security(
        device_id=device_id,
        assessment_id=assessment_id,
        custom_price=asm.get("price_lkr") if asm else None,
    )

    return {
        "assessment": asm,
        "device": dev,
        "findings": findings,
        "evaluation": evaluation,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def generate_markdown_report(assessment_id: str, output_path: str | None = None) -> str:
    """
    Generate Markdown format security assessment report with in-depth threat modeling.
    """
    data = get_assessment_report_data(assessment_id)
    asm = data["assessment"]
    dev = data["device"]
    eval_res = data["evaluation"]
    rec = eval_res["recommendation"]
    findings = data["findings"]

    out_file = output_path or f"{REPORTS_DIR}/{assessment_id}_security_report.md"
    Path(out_file).parent.mkdir(parents=True, exist_ok=True)

    lines = []
    lines.append("# WiFiRisk Security Assessment & Threat Modeling Report")
    lines.append("")
    lines.append("## Academic Project Context")
    lines.append("- **Project Title:** WiFiRisk: A Lightweight Security Assessment Framework for Low-Cost Wi-Fi Repeaters")
    lines.append("- **Student Name:** W.M.D.C.D.S Weerakoon")
    lines.append("- **Student ID:** 11161")
    lines.append("- **Faculty:** Faculty of Computer Science and Engineering, KIU")
    lines.append("- **Module:** COM4901 (Final Year Individual Research Project)")
    lines.append("")
    lines.append("---")
    lines.append("")
    lines.append("## 1. Assessment Overview")
    lines.append(f"- **Assessment ID:** `{asm.get('id')}`")
    lines.append(f"- **Assessment Status:** {asm.get('status')}")
    lines.append(f"- **Date Generated:** {data['generated_at']}")
    lines.append(f"- **Target Gateway IP:** `{asm.get('target_ip')}`")
    lines.append(f"- **Device Identifier:** `{asm.get('device_id')}`")
    lines.append(f"- **Device Brand / Model:** {dev.get('brand', 'Generic') if dev else 'Generic'} {dev.get('model', 'Wi-Fi Repeater') if dev else 'Wi-Fi Repeater'}")
    lines.append(f"- **Firmware Version:** `{dev.get('firmware_version', 'Unknown') if dev else 'Unknown'}`")
    lines.append(f"- **Purchase Price:** LKR {asm.get('price_lkr', 0)}")
    lines.append(f"- **Assessment Notes:** {asm.get('notes', 'None')}")
    lines.append("")
    lines.append("## 2. Executive Summary & Security Scorecard")
    lines.append(f"- **Security Score:** **{eval_res['final_score']} / 100**")
    lines.append(f"- **Base Points:** {eval_res['base_score']} (Total Deductions: -{eval_res['total_deductions']} points)")
    lines.append(f"- **Risk Level:** **{eval_res['risk_level']}**")
    lines.append(f"- **Price-to-Security Ratio (PSR):** **{eval_res['psr']}** (Baseline Premium: LKR 15,000)")
    lines.append(f"- **Recommendation Category:** **{rec['category']}**")
    lines.append("")
    lines.append(f"> **Guidance:** {rec['guidance']}")
    lines.append(">")
    lines.append(f"> **Price Evaluation:** {rec['price_verdict']}")
    lines.append("")
    lines.append("### Score Deductions Breakdown")
    lines.append("| Finding ID | Severity | Points Deducted | Title |")
    lines.append("| :--- | :---: | :---: | :--- |")
    for d in eval_res.get("deductions_detail", []):
        lines.append(f"| `{d.get('finding_id')}` | {d.get('severity')} | -{d.get('deduction')} | {d.get('title')} |")
    lines.append("")
    lines.append("## 3. Detailed Vulnerability Findings & Threat Modeling")
    if not findings:
        lines.append("*No security vulnerabilities or exposures were detected on the target device.*")
    else:
        for idx, f in enumerate(findings, 1):
            lines.append(f"### {idx}. [{f.get('id')}] {f.get('title')}")
            f_sev = f.get("severity") or f.get("level", "")
            lines.append(f"- **Severity:** {f_sev}")
            lines.append(f"- **Status:** {f.get('status', 'Confirmed')}")
            lines.append(f"- **Weakness Mapping:** {f.get('cwe', 'CWE-General Insecure Configuration')}")
            lines.append(f"- **Category:** {f.get('category')}")
            lines.append(f"- **Module:** {f.get('module', 'Assessment Engine')}")
            if f.get("evidence"):
                lines.append(f"- **Observed Technical Evidence:** `{f.get('evidence')}`")
            if f.get("threat_scenario"):
                lines.append(f"- **Threat Modeling & Attack Vector:** {f.get('threat_scenario')}")
            lines.append(f"- **Security Impact:** {f.get('impact')}")
            lines.append(f"- **Recommended Mitigation:** {f.get('recommendation')}")
            if f.get("hardening_coverage"):
                lines.append("- **Hardening Action Steps:**")
                for step in f.get("hardening_coverage", "").split("\n"):
                    if step.strip():
                        lines.append(f"  {step.strip()}")
            lines.append("")

    lines.append("## 4. Academic Disclaimer")
    lines.append("This security assessment report was generated automatically by WiFiRisk for academic research purposes under module COM4901. Findings reflect non-destructive protocol probes and vulnerability checks conducted in an isolated lab environment.")

    content = "\n".join(lines)
    with open(out_file, "w", encoding="utf-8") as f:
        f.write(content)

    return out_file


def generate_text_report(assessment_id: str, output_path: str | None = None) -> str:
    """
    Generate clean, monospaced ASCII plain-text (.txt) security report.
    """
    data = get_assessment_report_data(assessment_id)
    asm = data["assessment"]
    dev = data["device"]
    eval_res = data["evaluation"]
    rec = eval_res["recommendation"]
    findings = data["findings"]

    out_file = output_path or f"{REPORTS_DIR}/{assessment_id}_security_report.txt"
    Path(out_file).parent.mkdir(parents=True, exist_ok=True)

    lines = []
    lines.append("=" * 80)
    lines.append("WIFIRISK SECURITY ASSESSMENT & THREAT MODELING REPORT (ACADEMIC COM4901)")
    lines.append("=" * 80)
    lines.append(f"Assessment ID : {asm.get('id')}")
    lines.append(f"Target IP     : {asm.get('target_ip')}")
    lines.append(f"Device Model  : {dev.get('brand', 'Generic') if dev else 'Generic'} {dev.get('model', 'Wi-Fi Repeater') if dev else 'Wi-Fi Repeater'} (ID: {asm.get('device_id')})")
    lines.append(f"Firmware      : {dev.get('firmware_version', 'Unknown') if dev else 'Unknown'}")
    lines.append(f"Price (LKR)   : {asm.get('price_lkr', 0)}")
    lines.append(f"Date Exported : {data['generated_at']}")
    lines.append("-" * 80)
    lines.append("EXECUTIVE SECURITY EVALUATION & PSR SCORECARD")
    lines.append(f"Security Score : {eval_res['final_score']} / 100")
    lines.append(f"Risk Level     : {eval_res['risk_level']}")
    lines.append(f"PSR Ratio      : {eval_res['psr']} (Price-to-Security Ratio)")
    lines.append(f"Recommendation : {rec['category']}")
    lines.append(f"Guidance       : {rec['guidance']}")
    lines.append(f"Verdict        : {rec['price_verdict']}")
    lines.append("-" * 80)
    lines.append("SCORE DEDUCTIONS BREAKDOWN")
    for d in eval_res.get("deductions_detail", []):
        lines.append(f"  - [{d.get('finding_id')}] ({d.get('severity')}) -{d.get('deduction')} pts: {d.get('title')}")
    lines.append("-" * 80)
    lines.append(f"DETAILED FINDINGS & THREAT MODELING ({len(findings)})")
    if not findings:
        lines.append("  No security vulnerabilities detected.")
    for idx, f in enumerate(findings, 1):
        f_sev = f.get("severity") or f.get("level", "")
        lines.append(f"\n{idx}. [{f.get('id')}] {f.get('title')}")
        lines.append(f"   Severity      : {f_sev} | Status: {f.get('status', 'Confirmed')}")
        lines.append(f"   Weakness      : {f.get('cwe', 'CWE-General Insecure Configuration')}")
        lines.append(f"   Category      : {f.get('category')}")
        lines.append(f"   Module        : {f.get('module', 'Assessment Engine')}")
        if f.get("threat_scenario"):
            lines.append(f"   Threat Vector : {f.get('threat_scenario')}")
        if f.get("evidence"):
            lines.append(f"   Evidence      : {f.get('evidence')}")
        lines.append(f"   Impact        : {f.get('impact')}")
        lines.append(f"   Remediation   : {f.get('recommendation')}")
        if f.get("hardening_coverage"):
            lines.append("   Hardening     :")
            for h_step in f.get("hardening_coverage", "").splitlines():
                lines.append(f"     * {h_step}")
    lines.append("\n" + "-" * 80)
    lines.append("MITIGATION & HARDENING ACTION PLAN")
    for idx, m in enumerate(rec.get("mitigations", []), 1):
        lines.append(f"  {idx}. {m}")
    lines.append("=" * 80)

    content = "\n".join(lines)
    with open(out_file, "w", encoding="utf-8") as f:
        f.write(content)

    return out_file


def generate_docx_report(assessment_id: str, output_path: str | None = None) -> str:
    """
    Generate Microsoft Word (.docx) format security assessment report using python-docx.
    """
    import docx
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    data = get_assessment_report_data(assessment_id)
    asm = data["assessment"]
    dev = data["device"]
    eval_res = data["evaluation"]
    rec = eval_res["recommendation"]
    findings = data["findings"]

    out_file = output_path or f"{REPORTS_DIR}/{assessment_id}_security_report.docx"
    Path(out_file).parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()

    # Set document title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("WiFiRisk Security Assessment & Threat Modeling Report")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0, 51, 102)

    sub_p = doc.add_paragraph("A Lightweight Security Assessment Framework for Low-Cost Wi-Fi Repeaters")
    sub_p.runs[0].font.size = Pt(13)
    sub_p.runs[0].font.italic = True

    # Academic metadata
    doc.add_heading("Academic Project Context", level=2)
    p_meta = doc.add_paragraph()
    p_meta.add_run("Student Name: ").bold = True
    p_meta.add_run("W.M.D.C.D.S Weerakoon (ID: 11161)\n")
    p_meta.add_run("Faculty: ").bold = True
    p_meta.add_run("Faculty of Computer Science and Engineering, KIU\n")
    p_meta.add_run("Module: ").bold = True
    p_meta.add_run("COM4901 - Final Year Individual Research Project")

    # Assessment details table
    doc.add_heading("1. Assessment Details", level=2)
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    details_data = [
        ("Assessment ID", asm.get("id", "")),
        ("Date Generated", data["generated_at"]),
        ("Target IP / Gateway", asm.get("target_ip", "")),
        ("Device ID / Model", f"{asm.get('device_id', '')} ({dev.get('brand', 'Generic') if dev else 'Generic'} {dev.get('model', 'Wi-Fi Repeater') if dev else 'Wi-Fi Repeater'})"),
        ("Firmware Version", dev.get("firmware_version", "Unknown") if dev else "Unknown"),
        ("Purchase Price", f"LKR {asm.get('price_lkr', 0)}"),
    ]

    for row_idx, (k, v) in enumerate(details_data):
        row = table.rows[row_idx]
        row.cells[0].text = k
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[1].text = str(v)

    # Executive Scorecard
    doc.add_heading("2. Executive Summary & Security Scorecard", level=2)
    score_p = doc.add_paragraph()
    score_run = score_p.add_run(f"Security Score: {eval_res['final_score']} / 100  |  Risk Level: {eval_res['risk_level']}")
    score_run.bold = True
    score_run.font.size = Pt(14)
    if eval_res["risk_level"] in ["Critical", "High"]:
        score_run.font.color.rgb = RGBColor(180, 0, 0)
    else:
        score_run.font.color.rgb = RGBColor(0, 120, 0)

    doc.add_paragraph(f"Price-to-Security Ratio (PSR): {eval_res['psr']} (Baseline: LKR 15,000)")
    doc.add_paragraph(f"Recommendation Category: {rec['category']}").runs[0].bold = True
    doc.add_paragraph(f"Evaluation Guidance: {rec['guidance']}")
    doc.add_paragraph(f"Price Verdict: {rec['price_verdict']}")

    # Score deductions table
    doc.add_heading("Score Deductions Breakdown", level=3)
    ded_table = doc.add_table(rows=1, cols=4)
    ded_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = ded_table.rows[0].cells
    hdr[0].text = "Finding ID"
    hdr[1].text = "Severity"
    hdr[2].text = "Deduction"
    hdr[3].text = "Vulnerability Title"
    for cell in hdr:
        cell.paragraphs[0].runs[0].bold = True

    for d in eval_res.get("deductions_detail", []):
        row_cells = ded_table.add_row().cells
        row_cells[0].text = d.get("finding_id", "")
        row_cells[1].text = d.get("severity", "")
        row_cells[2].text = f"-{d.get('deduction', 0)}"
        row_cells[3].text = d.get("title", "")

    # Findings section with Threat Modeling
    doc.add_heading(f"3. Detailed Vulnerability Findings & Threat Modeling ({len(findings)})", level=2)
    for idx, f in enumerate(findings, 1):
        doc.add_heading(f"{idx}. [{f.get('id')}] {f.get('title')}", level=3)
        fp = doc.add_paragraph()
        f_sev = f.get("severity") or f.get("level", "")
        fp.add_run("Severity: ").bold = True
        fp.add_run(f"{f_sev}  |  ")
        fp.add_run("Status: ").bold = True
        fp.add_run(f"{f.get('status', 'Confirmed')}  |  ")
        fp.add_run("Weakness: ").bold = True
        fp.add_run(f"{f.get('cwe', 'CWE-General Insecure Configuration')}\n")

        if f.get("threat_scenario"):
            fp.add_run("Threat Modeling & Attack Vector: ").bold = True
            fp.add_run(f"{f.get('threat_scenario')}\n")

        if f.get("evidence"):
            fp.add_run("Observed Evidence: ").bold = True
            fp.add_run(f"{f.get('evidence')}\n")

        fp.add_run("Security Impact: ").bold = True
        fp.add_run(f"{f.get('impact')}\n")

        fp.add_run("Remediation: ").bold = True
        fp.add_run(f"{f.get('recommendation')}\n")

        if f.get("hardening_coverage"):
            fp.add_run("Hardening Actions:\n").bold = True
            for h_step in f.get("hardening_coverage", "").splitlines():
                fp.add_run(f"  • {h_step}\n")

    # Mitigation section
    doc.add_heading("4. Recommended Hardening Actions", level=2)
    for idx, m in enumerate(rec.get("mitigations", []), 1):
        doc.add_paragraph(f"{idx}. {m}")

    doc.save(out_file)
    return out_file


def export_assessment_report(
    assessment_id: str,
    format_type: str = "all",
    output_dir: str = REPORTS_DIR,
) -> dict[str, str]:
    """
    Export reports in Markdown, Plain Text, DOCX or all formats.
    """
    Path(output_dir).mkdir(parents=True, exist_ok=True)
    generated_files: dict[str, str] = {}

    fmt = format_type.lower().strip()

    if fmt in ["md", "markdown", "all"]:
        md_path = f"{output_dir}/{assessment_id}_security_report.md"
        generate_markdown_report(assessment_id, md_path)
        generated_files["markdown"] = md_path

    if fmt in ["txt", "text", "all"]:
        txt_path = f"{output_dir}/{assessment_id}_security_report.txt"
        generate_text_report(assessment_id, txt_path)
        generated_files["text"] = txt_path

    if fmt in ["docx", "word", "all"]:
        docx_path = f"{output_dir}/{assessment_id}_security_report.docx"
        generate_docx_report(assessment_id, docx_path)
        generated_files["docx"] = docx_path

    return generated_files


def generate_research_dataset_txt(output_path: str | None = None) -> str:
    """
    Generate a comprehensive academic research test dataset in plain text (.txt) format,
    aggregating all tested devices, assessment sessions, vulnerability findings, scores,
    and technical evidence logs for research supervisors and review boards.
    """
    from wifi_risk.services.assessment_service import get_all_assessments
    from wifi_risk.services.device_service import get_all_devices
    from wifi_risk.services.finding_service import (
        get_all_findings,
        get_findings_by_assessment_id,
        get_findings_by_device_id,
    )
    from wifi_risk.services.scoring_service import evaluate_device_security

    out_file = output_path or f"{REPORTS_DIR}/wifi_risk_research_dataset.txt"
    Path(out_file).parent.mkdir(parents=True, exist_ok=True)

    assessments = get_all_assessments()
    devices = get_all_devices()
    all_findings = get_all_findings()
    gen_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    total_assessments = len(assessments)
    total_devices = len(devices)
    total_findings_count = len(all_findings)
    high_count = sum(1 for f in all_findings if (f.get("severity") or f.get("level", "")).lower() in ["high", "critical"])
    med_count = sum(1 for f in all_findings if "medium" in (f.get("severity") or f.get("level", "")).lower())
    low_med_count = sum(1 for f in all_findings if (f.get("severity") or f.get("level", "")).lower() == "low-medium")
    low_count = sum(1 for f in all_findings if (f.get("severity") or f.get("level", "")).lower() == "low")
    info_count = sum(1 for f in all_findings if (f.get("severity") or f.get("level", "")).lower() == "informational")

    lines = []
    lines.append("=" * 115)
    lines.append("WIFIRISK ACADEMIC RESEARCH TEST DATASET & VULNERABILITY AUDIT RECORD")
    lines.append("=" * 115)
    lines.append("ACADEMIC PROJECT METADATA:")
    lines.append("  • Student Name     : W.M.D.C.D.S Weerakoon (Student ID: 11161)")
    lines.append("  • Degree / Faculty : Faculty of Computer Science and Engineering, KIU")
    lines.append("  • Research Module  : COM4901 (Final Year Individual Research Project)")
    lines.append("  • Research Title   : Security Assessment of Low-Cost Wi-Fi Repeaters: A Study of Consumer Devices")
    lines.append(f"  • Export Date/Time : {gen_time}")
    lines.append(f"  • Total Scope      : {total_assessments} Assessment Sessions | {total_devices} Cataloged Devices | {total_findings_count} Security Findings")
    lines.append("=" * 115)
    lines.append("")

    lines.append("SECTION 1: DATASET AGGREGATE SUMMARY & SEVERITY DISTRIBUTION")
    lines.append("-" * 115)
    lines.append(f"Total Assessment Sessions Logged : {total_assessments}")
    lines.append(f"Total Tested Repeaters / Devices : {total_devices}")
    lines.append(f"Total Confirmed Vulnerabilities  : {total_findings_count}")
    lines.append(f"  ├── High / Critical Severity   : {high_count} finding(s)")
    lines.append(f"  ├── Medium Severity            : {med_count} finding(s)")
    lines.append(f"  ├── Low-Medium Severity        : {low_med_count} finding(s)")
    lines.append(f"  ├── Low Severity               : {low_count} finding(s)")
    lines.append(f"  └── Informational Notices      : {info_count} finding(s)")
    lines.append("-" * 115)
    lines.append("")

    lines.append("SECTION 2: MASTER DEVICE COMPARATIVE BENCHMARK MATRIX")
    lines.append("-" * 115)
    lines.append(f"{'Device ID':<10} | {'Brand & Model':<28} | {'Price (LKR)':<11} | {'Score':<6} | {'Risk Level':<9} | {'PSR':<5} | {'Findings (H/M/L)':<16} | {'Recommendation':<20}")
    lines.append("-" * 115)

    for dev in devices:
        dev_id = dev.get("id", "")
        brand_model = f"{dev.get('brand', '')} {dev.get('model', '')}".strip()[:28]
        price = dev.get("price_lkr", 0)
        score = dev.get("security_score", 0)
        risk = dev.get("risk_level", "")
        rec = dev.get("recommendation", "")[:20]

        psr_calc = round((score / 100.0) / (price / 15000.0), 2) if price > 0 else 0.0
        dev_findings = get_findings_by_device_id(dev_id)
        d_high = sum(1 for f in dev_findings if (f.get("severity") or f.get("level", "")).lower() in ["high", "critical"])
        d_med = sum(1 for f in dev_findings if "medium" in (f.get("severity") or f.get("level", "")).lower())
        d_low = len(dev_findings) - d_high - d_med
        findings_breakdown = f"{len(dev_findings)} ({d_high}H/{d_med}M/{d_low}L)"

        lines.append(f"{dev_id:<10} | {brand_model:<28} | {price:<11} | {score:<6} | {risk:<9} | {psr_calc:<5.2f} | {findings_breakdown:<16} | {rec:<20}")

    lines.append("-" * 115)
    lines.append("")

    lines.append("SECTION 3: DETAILED ASSESSMENT SESSIONS & TEST EXECUTION LOGS")
    lines.append("=" * 115)

    for asm in assessments:
        asm_id = asm.get("id", "")
        asm_dev_id = asm.get("device_id", "")
        target_ip = asm.get("target_ip", "")
        price_lkr = asm.get("price_lkr", 0)
        status = asm.get("status", "Created")
        created_at = asm.get("created_at", "")
        notes = asm.get("notes", "None")
        checks = asm.get("checks", [])

        eval_res = evaluate_device_security(
            device_id=asm_dev_id if not asm_id else None,
            assessment_id=asm_id,
            custom_price=price_lkr,
        )
        asm_findings = get_findings_by_assessment_id(asm_id)

        lines.append(f"\n[ASSESSMENT SESSION: {asm_id}]")
        lines.append("-" * 75)
        lines.append(f"  • Device Identifier  : {asm_dev_id}")
        lines.append(f"  • Target Gateway IP  : {target_ip}")
        lines.append(f"  • Purchase Price     : LKR {price_lkr}")
        lines.append(f"  • Assessment Status  : {status}")
        lines.append(f"  • Session Timestamp  : {created_at}")
        lines.append(f"  • Session Notes      : {notes}")
        lines.append(f"  • Final Score & Tier : {eval_res['final_score']}/100 ({eval_res['risk_level']} Risk)")
        lines.append(f"  • Price-to-Security  : PSR {eval_res['psr']} (Baseline: LKR 15,000)")
        lines.append(f"  • Verdict Guidance   : {eval_res['recommendation'].get('guidance', '')}")

        if checks:
            lines.append("\n  Audit Modules Executed:")
            for c in checks:
                mod = c.get("module", "Check")
                ts = c.get("timestamp", "")
                f_count = c.get("findings_count", 0)
                ports = c.get("open_ports", [])
                ports_str = f" | Open Ports: {ports}" if ports else ""
                lines.append(f"    - [{mod}] Timestamp: {ts} | Findings Detected: {f_count}{ports_str}")

        if asm_findings:
            lines.append(f"\n  Confirmed Vulnerabilities Discovered ({len(asm_findings)} total):")
            for idx, f in enumerate(asm_findings, 1):
                f_id = f.get("id", "")
                f_sev = f.get("severity") or f.get("level", "")
                f_title = f.get("title", "")
                f_module = f.get("module", "General")
                f_cwe = f.get("cwe", "CWE-General")
                f_ev = f.get("evidence", "").strip()
                f_impact = f.get("impact", "")
                f_rec = f.get("recommendation", "")

                lines.append(f"\n    {idx}. [{f_id}] [{f_sev.upper()}] {f_title}")
                lines.append(f"       • Layer / Module    : {f_module}")
                lines.append(f"       • MITRE Weakness    : {f_cwe}")
                if f_ev:
                    lines.append(f"       • Observed Evidence : {f_ev}")
                lines.append(f"       • Threat Impact     : {f_impact}")
                lines.append(f"       • Remediation Plan  : {f_rec}")
        else:
            lines.append("\n  No security findings recorded for this session (Clean baseline or pending tests).")

        lines.append("." * 75)

    lines.append("\n" + "=" * 115)
    lines.append("END OF RESEARCH TEST DATASET RECORD")
    lines.append("=" * 115)

    content = "\n".join(lines)
    with open(out_file, "w", encoding="utf-8") as f:
        f.write(content)

    return out_file


def generate_research_audit_log(output_path: str | None = None) -> str:
    """
    Generate a standardized, timestamped chronological audit log (.log) format of all
    assessments, network probes, web scans, and security findings.
    """
    from wifi_risk.services.assessment_service import get_all_assessments
    from wifi_risk.services.finding_service import get_all_findings

    out_file = output_path or f"{REPORTS_DIR}/wifi_risk_audit_log.log"
    Path(out_file).parent.mkdir(parents=True, exist_ok=True)

    assessments = get_all_assessments()
    findings = get_all_findings()
    log_lines = []

    log_lines.append("# WiFiRisk Automated Security Audit Trail Log - COM4901 Research")
    log_lines.append(f"# Generated: {datetime.now().isoformat(timespec='seconds')}")
    log_lines.append("# Target Research Scope: Low-Cost Wi-Fi Repeaters & Consumer IoT Devices")
    log_lines.append("# Format: [TIMESTAMP] [SEVERITY/EVENT] [ASSESSMENT_ID] [DEVICE_ID] MESSAGE")
    log_lines.append("# " + "-" * 90)

    for asm in assessments:
        asm_id = asm.get("id", "")
        dev_id = asm.get("device_id", "")
        ts = asm.get("created_at", datetime.now().isoformat(timespec="seconds"))
        target_ip = asm.get("target_ip", "")
        price = asm.get("price_lkr", 0)

        log_lines.append(f"[{ts}] [SESSION_INIT] [{asm_id}] [{dev_id}] Assessment session initialized on target {target_ip} (Price: LKR {price})")

        checks = asm.get("checks", [])
        for c in checks:
            c_ts = c.get("timestamp", ts)
            c_mod = c.get("module", "Module")
            c_fcount = c.get("findings_count", 0)
            c_ports = c.get("open_ports", [])
            log_lines.append(f"[{c_ts}] [MODULE_EXEC] [{asm_id}] [{dev_id}] {c_mod} completed. Discovered: {c_fcount} finding(s), Open Ports: {c_ports}")

    for f in findings:
        f_id = f.get("id", "")
        f_asm = f.get("assessment_id", "GLOBAL")
        f_dev = f.get("device_id", "DEVICE")
        f_sev = (f.get("severity") or f.get("level", "MEDIUM")).upper()
        f_title = f.get("title", "")
        f_mod = f.get("module", "General")
        f_ts = f.get("date_observed", datetime.now().isoformat(timespec="seconds"))
        f_cwe = f.get("cwe", "CWE-General")

        log_lines.append(f"[{f_ts}] [VULN_ALERT] [{f_asm}] [{f_dev}] [{f_sev}] [{f_id}] {f_title} ({f_mod} | {f_cwe})")

    content = "\n".join(log_lines)
    with open(out_file, "w", encoding="utf-8") as f:
        f.write(content)

    return out_file


def export_research_test_dataset(
    output_dir: str = REPORTS_DIR,
    format_type: str = "all",
) -> dict[str, str]:
    """
    Export the full research test dataset and audit logs in TXT, LOG, JSON or all formats.
    """
    Path(output_dir).mkdir(parents=True, exist_ok=True)
    generated_files: dict[str, str] = {}
    fmt = format_type.lower().strip()

    if fmt in ["txt", "text", "all"]:
        txt_path = f"{output_dir}/wifi_risk_research_dataset.txt"
        generate_research_dataset_txt(txt_path)
        generated_files["dataset_txt"] = txt_path

    if fmt in ["log", "audit", "all"]:
        log_path = f"{output_dir}/wifi_risk_audit_log.log"
        generate_research_audit_log(log_path)
        generated_files["audit_log"] = log_path

    if fmt in ["json", "data", "all"]:
        import json
        from wifi_risk.services.assessment_service import get_all_assessments
        from wifi_risk.services.device_service import get_all_devices
        from wifi_risk.services.finding_service import get_all_findings

        json_path = f"{output_dir}/wifi_risk_dataset.json"
        master_json = {
            "metadata": {
                "project": "COM4901 Final Year Research Project",
                "student": "W.M.D.C.D.S Weerakoon (ID: 11161)",
                "faculty": "Faculty of Computer Science and Engineering, KIU",
                "exported_at": datetime.now().isoformat(timespec="seconds"),
            },
            "devices": get_all_devices(),
            "assessments": get_all_assessments(),
            "findings": get_all_findings(),
        }
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(master_json, f, indent=2)
        generated_files["dataset_json"] = json_path

    return generated_files

